import os
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ALLOWED_EXTENSIONS = {
    ".py", ".js", ".ts", ".java", ".cpp", ".c", ".h",
    ".html", ".css", ".json", ".md"
}

IGNORED_DIRS = {
    ".git", "__pycache__", "node_modules",
    "dist", "build", ".idea", ".vscode"
}

OUTPUT_FILE = "listing.docx"


def should_include(file):
    return any(file.endswith(ext) for ext in ALLOWED_EXTENSIONS)


def set_two_columns(section):
    sectPr = section._sectPr

    cols = OxmlElement('w:cols')
    cols.set(qn('w:num'), '2')
    sectPr.append(cols)


def add_code(doc, text_lines):
    code = "".join(text_lines)

    p = doc.add_paragraph()
    run = p.add_run(code)

    font = run.font
    font.name = "Courier New"
    font.size = Pt(8)


def read_file(path):
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.readlines()
    except:
        return None


def main():
    doc = Document()

    set_two_columns(doc.sections[0])

    doc.add_heading("Project Code Listing (2 columns)", 0)

    for root, dirs, files in os.walk("."):
        dirs[:] = [d for d in dirs if d not in IGNORED_DIRS]

        for file in files:
            if not should_include(file):
                continue

            path = os.path.join(root, file)
            lines = read_file(path)

            if not lines:
                continue

            doc.add_heading(path, level=2)
            doc.add_paragraph(f"Lines: {len(lines)}")
            add_code(doc, lines)

    doc.save(OUTPUT_FILE)
    print("Готово:", OUTPUT_FILE)


if __name__ == "__main__":
    main()